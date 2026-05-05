import csv
import io
from decimal import Decimal, InvalidOperation

from django.utils.text import slugify
from openpyxl import load_workbook

from products.models import Category, Product

from . import importers_prom

ALLOWED_LANGUAGES = ('uk', 'ru')
DEFAULT_LANGUAGE = 'uk'


def _normalize_language(language):
    return language if language in ALLOWED_LANGUAGES else DEFAULT_LANGUAGE

_COLUMN_ALIASES = {
    # sku
    'код товару': 'sku',
    'код': 'sku',
    'артикул': 'sku',
    'sku': 'sku',
    # name
    'назва': 'name',
    'назва товару': 'name',
    'найменування': 'name',
    'name': 'name',
    # sku_manufacturer
    'sku виробника': 'sku_manufacturer',
    'артикул виробника': 'sku_manufacturer',
    'sku_manufacturer': 'sku_manufacturer',
    # category
    'категорія': 'category',
    'категория': 'category',
    'category': 'category',
    # price
    'ціна': 'price',
    'цена': 'price',
    'price': 'price',
    # old_price
    'стара ціна': 'old_price',
    'стара цена': 'old_price',
    'old_price': 'old_price',
    # stock_status
    'наявність': 'stock_status',
    'статус': 'stock_status',
    'stock_status': 'stock_status',
    # badge
    'бейдж': 'badge',
    'badge': 'badge',
    # discount_percent
    'знижка %': 'discount_percent',
    'знижка': 'discount_percent',
    'discount_percent': 'discount_percent',
    # manufacturer
    'виробник': 'manufacturer',
    'manufacturer': 'manufacturer',
    # country
    'країна': 'country',
    'country': 'country',
    # weight
    'вага': 'weight',
    'weight': 'weight',
    # condition
    'стан': 'condition',
    'condition': 'condition',
    # description
    'опис': 'description',
    'опис (html)': 'description',
    'description': 'description',
    # short_description
    'короткий опис': 'short_description',
    'short_description': 'short_description',
    # meta_title
    'meta title': 'meta_title',
    'мета заголовок': 'meta_title',
    'meta_title': 'meta_title',
    # meta_description
    'meta description': 'meta_description',
    'мета опис': 'meta_description',
    'meta_description': 'meta_description',
    # is_active
    'активний': 'is_active',
    'активно': 'is_active',
    'is_active': 'is_active',
    # id (read-only, not used for create/update)
    'id': 'id',
}


def _normalize_row(row_dict):
    """Return a new dict with keys normalized to canonical English field names."""
    normalized = {}
    for key, value in row_dict.items():
        canonical = _COLUMN_ALIASES.get(key.strip().lower(), key.strip().lower())
        normalized[canonical] = value
    return normalized


def _process_row(row_dict):
    """row_dict must already be normalized (keys are canonical English field names)."""
    sku = row_dict.get('sku', '').strip()
    name = row_dict.get('name', '').strip()
    if not sku or not name:
        return None, 'Відсутній SKU або назва'

    raw_categories = row_dict.get('category', '').strip()
    category_objs = []
    for cat_name in (c.strip() for c in raw_categories.replace(';', ',').split(',') if c.strip()):
        cat, _ = Category.objects.get_or_create(
            name=cat_name,
            defaults={'slug': slugify(cat_name, allow_unicode=True) or sku.lower()}
        )
        category_objs.append(cat)

    try:
        price = Decimal(row_dict.get('price', '0'))
    except (InvalidOperation, ValueError):
        price = Decimal('0')

    old_price = None
    old_price_str = row_dict.get('old_price', '').strip()
    if old_price_str:
        try:
            old_price = Decimal(old_price_str)
        except (InvalidOperation, ValueError):
            pass

    weight = None
    weight_str = row_dict.get('weight', '').strip()
    if weight_str:
        try:
            weight = Decimal(weight_str)
        except (InvalidOperation, ValueError):
            pass

    base_slug = slugify(name, allow_unicode=True)[:400] or sku.lower()
    slug = base_slug
    counter = 1
    while Product.objects.filter(slug=slug).exclude(sku=sku).exists():
        slug = f"{base_slug[:395]}-{counter}"
        counter += 1

    try:
        discount_pct = int(row_dict.get('discount_percent', 0) or 0)
    except (ValueError, TypeError):
        discount_pct = 0

    defaults = {
        'name': name,
        'slug': slug,
        'price': price,
        'old_price': old_price,
        'stock_status': row_dict.get('stock_status', 'in_stock').strip() or 'in_stock',
        'badge': row_dict.get('badge', '').strip(),
        'discount_percent': discount_pct,
        'manufacturer': row_dict.get('manufacturer', '').strip(),
        'country': row_dict.get('country', '').strip(),
        'weight': weight,
        'condition': row_dict.get('condition', 'Новий').strip() or 'Новий',
        'sku_manufacturer': row_dict.get('sku_manufacturer', '').strip(),
        'description': row_dict.get('description', '').strip(),
        'short_description': row_dict.get('short_description', '').strip(),
        'meta_title': row_dict.get('meta_title', '').strip(),
        'meta_description': row_dict.get('meta_description', '').strip(),
        'is_active': str(row_dict.get('is_active', 'True')).lower() in ('true', '1', 'yes'),
    }

    product, created = Product.objects.update_or_create(sku=sku, defaults=defaults)
    if category_objs:
        product.categories.set(category_objs)
    return product, 'створено' if created else 'оновлено'


def _row_to_str_dict(headers, row):
    return {k: ('' if v is None else str(v).strip()) for k, v in zip(headers, row)}


def _process_one(headers_raw, headers_lower, row_values, *, fmt, language, fetch_images):
    if fmt == 'prom':
        row_dict = importers_prom.build_prom_row(headers_lower, row_values, language)
        attributes = importers_prom.parse_prom_attributes(headers_lower, row_values)
        image_urls = importers_prom.parse_image_urls(row_dict.pop('_image_urls', ''))
        product, status = importers_prom.process_prom_row(
            row_dict, attributes, image_urls, fetch_images=fetch_images,
        )
        sku = row_dict.get('sku', '')
    else:
        raw = _row_to_str_dict(headers_raw, row_values)
        row_dict = _normalize_row(raw)
        product, status = _process_row(row_dict)
        sku = row_dict.get('sku', '')
    return sku, status


def _import_rows(headers_raw, body_rows, *, language, fetch_images):
    headers_lower = [str(h).strip().lower() for h in headers_raw]
    fmt = 'prom' if importers_prom.is_prom_format(headers_lower) else 'legacy'
    results = []
    for i, row in enumerate(body_rows, start=2):
        try:
            sku, status = _process_one(
                headers_raw, headers_lower, row,
                fmt=fmt, language=language, fetch_images=fetch_images,
            )
        except Exception as exc:
            sku = ''
            status = f'Помилка: {exc}'
        results.append({'row': i, 'sku': sku, 'status': status})
    return results


def import_csv(file_obj, *, language=DEFAULT_LANGUAGE, fetch_images=True):
    language = _normalize_language(language)
    content = file_obj.read().decode('utf-8-sig')
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)
    if not rows:
        return []
    headers_raw = [str(h).strip() for h in rows[0]]
    return _import_rows(headers_raw, rows[1:], language=language, fetch_images=fetch_images)


def import_excel(file_obj, *, language=DEFAULT_LANGUAGE, fetch_images=True):
    language = _normalize_language(language)
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb.active
    headers_raw = [str(cell.value).strip() if cell.value is not None else '' for cell in ws[1]]
    body = list(ws.iter_rows(min_row=2, values_only=True))
    return _import_rows(headers_raw, body, language=language, fetch_images=fetch_images)
